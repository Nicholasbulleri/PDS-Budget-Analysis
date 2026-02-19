#!/usr/bin/env python3
"""
Run closed_projects_missing_budget_by_business_line.sql: count of distinct closed USD projects
missing budgets in generictask, grouped by business line. Outputs CSV and Excel.
"""
import csv
import os

try:
    from dotenv import load_dotenv
    try:
        load_dotenv()
    except (PermissionError, FileNotFoundError):
        pass
except ImportError:
    pass

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
SQL_FILE = os.path.join(SCRIPT_DIR, "closed_projects_missing_budget_by_business_line.sql")
OUT_CSV = os.path.join(SCRIPT_DIR, "closed_projects_missing_budget_by_business_line.csv")
OUT_XLSX = os.path.join(SCRIPT_DIR, "closed_projects_missing_budget_by_business_line.xlsx")


def get_query():
    if os.path.isfile(SQL_FILE):
        with open(SQL_FILE, "r", encoding="utf-8") as f:
            return f.read().strip()
    raise FileNotFoundError(SQL_FILE)


def main():
    from edp_connection import execute_query

    query = get_query()
    print("Running closed projects missing budget by business line...")
    columns, results = execute_query(query)
    print(f"Retrieved {len(results)} rows, columns: {columns}")

    with open(OUT_CSV, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(columns)
        w.writerows(results)
    print(f"Results written to {OUT_CSV}")

    try:
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Missing budget by business line"
        for c, col in enumerate(columns, 1):
            ws.cell(row=1, column=c, value=col)
        for r, row in enumerate(results, 2):
            for c, val in enumerate(row, 1):
                ws.cell(row=r, column=c, value=val)
        wb.save(OUT_XLSX)
        print(f"Results written to {OUT_XLSX}")
    except ImportError:
        import pandas as pd
        pd.DataFrame(results, columns=columns).to_excel(
            OUT_XLSX, index=False, sheet_name="Missing budget by business line"
        )
        print(f"Results written to {OUT_XLSX} (via pandas)")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
