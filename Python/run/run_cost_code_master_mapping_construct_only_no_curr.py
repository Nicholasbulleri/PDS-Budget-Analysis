#!/usr/bin/env python3
"""Run cost_code_master_mapping_construct_only_no_curr.sql and write results to CSV + XLSX."""
import csv
import os

try:
    from dotenv import load_dotenv
    load_dotenv()
except ImportError:
    pass

import sys
from pathlib import Path

for _p in Path(__file__).resolve().parents:
    if (_p / "repo_paths.py").is_file():
        if str(_p) not in sys.path:
            sys.path.insert(0, str(_p))
        break
else:
    raise RuntimeError("Not inside project Python tree")
from repo_paths import repo_root

PROJECT_ROOT = repo_root()
SQL_DIR = os.path.join(PROJECT_ROOT, "sql", "cost_codes")
DATA_DIR = os.path.join(PROJECT_ROOT, "data", "cost_codes")
SQL_FILE = os.path.join(SQL_DIR, "cost_code_master_mapping_construct_only_no_curr.sql")
OUT_CSV = os.path.join(PROJECT_ROOT, "data", "cost_codes", "cost_code_master_mapping_construct_only_no_curr.csv")
OUT_XLSX = os.path.join(PROJECT_ROOT, "data", "cost_codes", "cost_code_master_mapping_construct_only_no_curr.xlsx")


def main():
    from edp_connection import execute_query

    with open(SQL_FILE, "r", encoding="utf-8") as f:
        query = f.read()

    print("Running cost code master mapping (construct only, no currency)...")
    columns, results = execute_query(query)
    print(f"Retrieved {len(results)} cost code mappings")

    os.makedirs(DATA_DIR, exist_ok=True)
    with open(OUT_CSV, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(columns)
        w.writerows(results)
    print(f"Written to {OUT_CSV}")

    try:
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Cost Code Mapping (Construct)"[:31]
        for c, col in enumerate(columns, 1):
            ws.cell(row=1, column=c, value=col)
        for r, row in enumerate(results, 2):
            for c, val in enumerate(row, 1):
                ws.cell(row=r, column=c, value=val)
        wb.save(OUT_XLSX)
        print(f"Written to {OUT_XLSX}")
    except ImportError:
        try:
            import pandas as pd
            pd.DataFrame(results, columns=columns).to_excel(OUT_XLSX, index=False, sheet_name="Cost Code Mapping"[:31])
            print(f"Written to {OUT_XLSX} (via pandas)")
        except Exception:
            pass

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
