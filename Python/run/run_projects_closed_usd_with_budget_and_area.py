#!/usr/bin/env python3
"""
Run projects query: distinct projects (Closed + Construction USD, with budget in generictask, with positive area).
- Uses projects_closed_usd_with_budget_and_area.sql
- Applies city->state mapping; writes CSV + Excel.
"""
import csv
import json
import os

try:
    from dotenv import load_dotenv
    try:
        load_dotenv()
    except (PermissionError, FileNotFoundError):
        pass
except ImportError:
    pass

import sys
from pathlib import Path

for _p in Path(__file__).resolve().parents:
    if (_p / "repo_paths.py").is_file():
        if str(_p) not in sys.path:
            sys.path.insert(0, str(_p))
        break
else:
    raise RuntimeError("Not inside project Python tree")
from repo_paths import repo_root

PROJECT_ROOT = repo_root()
SQL_DIR = os.path.join(PROJECT_ROOT, "sql", "projects")
DATA_DIR = os.path.join(PROJECT_ROOT, "data", "projects")
SQL_FILE = os.path.join(SQL_DIR, "projects_closed_usd_with_budget_and_area.sql")
OUT_CSV = os.path.join(PROJECT_ROOT, "data", "projects", "projects_closed_usd_with_budget_and_area.csv")
OUT_XLSX = os.path.join(PROJECT_ROOT, "data", "projects", "projects_closed_usd_with_budget_and_area.xlsx")

MAPPING_NAMES = [
    "city_to_state_mapping.json",
    "city_to_state_mapping.xlsx",
    "city_to_state_mapping.xls",
    "city_to_state_mapping.csv",
]


def normalize_key(city, country):
    c = (city or "").strip().lower()
    co = (country or "").strip().lower()
    return (c, co)


def load_mapping_from_json(path):
    with open(path, "r", encoding="utf-8") as f:
        data = json.load(f)
    if isinstance(data, list):
        return [{"city": (x.get("city") or "").strip(), "country": (x.get("country") or "").strip(), "state_abbrev": (x.get("state_abbrev") or x.get("state") or "").strip()} for x in data]
    if isinstance(data, dict):
        if "cities" in data or "mappings" in data:
            arr = data.get("cities", data.get("mappings", []))
            return [{"city": (x.get("city") or "").strip(), "country": (x.get("country") or "").strip(), "state_abbrev": (x.get("state_abbrev") or x.get("state") or "").strip()} for x in arr]
        out = []
        for k, v in data.items():
            if "|" in k:
                city, country = k.split("|", 1)
                out.append({"city": city.strip(), "country": country.strip(), "state_abbrev": (v or "").strip()})
        return out if out else []
    return []


def load_mapping_from_csv(path):
    rows = []
    with open(path, "r", encoding="utf-8") as f:
        r = csv.DictReader(f)
        for row in r:
            rows.append({"city": (row.get("city") or "").strip(), "country": (row.get("country") or "").strip(), "state_abbrev": (row.get("state_abbrev") or "").strip()})
    return rows


def load_mapping_from_excel(path):
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = []
        header = None
        for r in ws.iter_rows(values_only=True):
            if header is None:
                header = [str(c).strip().lower() if c else "" for c in r]
                ci = header.index("city") if "city" in header else 0
                co = header.index("country") if "country" in header else 1
                st = next((i for i, h in enumerate(header) if "state" in h), 2)
                continue
            cells = list(r) if r else []
            city = str(cells[ci]) if ci < len(cells) and cells[ci] is not None else ""
            country = str(cells[co]) if co < len(cells) and cells[co] is not None else ""
            state = str(cells[st]) if st < len(cells) and cells[st] is not None else ""
            rows.append({"city": city.strip(), "country": country.strip(), "state_abbrev": state.strip()})
        wb.close()
        return rows
    except ImportError:
        import pandas as pd
        df = pd.read_excel(path)
        df.columns = [str(c).strip().lower() for c in df.columns]
        city_col = "city" if "city" in df.columns else df.columns[0]
        country_col = "country" if "country" in df.columns else (df.columns[1] if len(df.columns) > 1 else "")
        state_col = next((c for c in df.columns if "state" in c), df.columns[2] if len(df.columns) > 2 else "")
        return [{"city": str(row.get(city_col, "")).strip(), "country": str(row.get(country_col, "")).strip(), "state_abbrev": str(row.get(state_col, "")).strip()} for _, row in df.iterrows()]


def build_lookup(rows):
    lookup = {}
    for r in rows:
        key = normalize_key(r["city"], r["country"])
        if key not in lookup:
            lookup[key] = r["state_abbrev"] or ""
    return lookup


# Country values that should be treated as US for state lookup (mapping file uses "United States")
US_COUNTRY_ALIASES = frozenset({"united states", "usa", "us", "u.s.", "u.s.a.", "united states of america"})


def find_and_load_mapping():
    for name in MAPPING_NAMES:
        path = os.path.join(DATA_DIR, name)
        if not os.path.isfile(path):
            continue
        ext = os.path.splitext(name)[1].lower()
        if ext == ".json":
            rows = load_mapping_from_json(path)
        elif ext in (".xlsx", ".xls"):
            rows = load_mapping_from_excel(path)
        else:
            rows = load_mapping_from_csv(path)
        lookup = build_lookup(rows)
        print(f"Loaded {len(rows)} mapping rows from {name} -> {len(lookup)} unique city/country keys.")
        return lookup
    print("No city->state mapping file found; state column will be empty.")
    return {}


def get_query():
    if os.path.isfile(SQL_FILE):
        with open(SQL_FILE, "r", encoding="utf-8") as f:
            return f.read().strip()
    raise FileNotFoundError(SQL_FILE)


def main():
    from edp_connection import execute_query

    lookup = find_and_load_mapping()
    query = get_query()
    print("Running projects query (Closed + Construction USD, budget + area)...")
    columns, results = execute_query(query)
    print(f"Retrieved {len(results)} project rows, columns: {columns}")

    out_columns = []
    for c in columns:
        out_columns.append(c)
        if c == "city":
            out_columns.append("state")
    if "state" not in out_columns:
        out_columns.append("state")

    ci = columns.index("city") if "city" in columns else -1
    co = columns.index("country") if "country" in columns else -1

    out_rows = []
    for row in results:
        row_list = list(row)
        city_val = row_list[ci] if ci >= 0 and ci < len(row_list) else None
        country_val = row_list[co] if co >= 0 and co < len(row_list) else None
        country_for_lookup = country_val if (country_val and str(country_val).strip()) else "United States"
        key = normalize_key(city_val, country_for_lookup)
        state_val = lookup.get(key, "")
        # If no match and country looks like US, try "United States" (mapping file uses this)
        if state_val == "" and (country_for_lookup or "").strip().lower() in US_COUNTRY_ALIASES:
            state_val = lookup.get(normalize_key(city_val, "United States"), "")
        if state_val == "" and country_for_lookup != (country_val or ""):
            state_val = lookup.get(normalize_key(city_val, country_val or ""), "")

        new_row = []
        for i, c in enumerate(columns):
            new_row.append(row_list[i] if i < len(row_list) else None)
            if c == "city":
                new_row.append(state_val or None)
        out_rows.append(new_row)

    with open(OUT_CSV, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(out_columns)
        w.writerows(out_rows)
    print(f"Results written to {OUT_CSV}")

    try:
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Projects"[:31]
        for c, col in enumerate(out_columns, 1):
            ws.cell(row=1, column=c, value=col)
        for r, row in enumerate(out_rows, 2):
            for c, val in enumerate(row, 1):
                ws.cell(row=r, column=c, value=val)
        wb.save(OUT_XLSX)
        print(f"Results written to {OUT_XLSX}")
    except ImportError:
        import pandas as pd
        pd.DataFrame(out_rows, columns=out_columns).to_excel(OUT_XLSX, index=False, sheet_name="Projects"[:31])
        print(f"Results written to {OUT_XLSX} (via pandas)")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
