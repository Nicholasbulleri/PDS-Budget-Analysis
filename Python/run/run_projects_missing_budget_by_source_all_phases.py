#!/usr/bin/env python3
"""
Run projects_missing_budget_by_source_all_phases.sql: projects missing budget by phase and source
(Clarizen, Ingenious) for the Missing Budget heatmap. Outputs CSV (and Excel) to data/.

From project root:
  python3 Python/run/run_projects_missing_budget_by_source_all_phases.py
  OR  ./run_missing_budget_heatmap.sh
From Python/run/ directory:
  python3 run_projects_missing_budget_by_source_all_phases.py
"""
import csv
import os

try:
    from dotenv import load_dotenv
    try:
        load_dotenv()
    except (PermissionError, FileNotFoundError):
        pass
except ImportError:
    pass

import sys
from pathlib import Path

for _p in Path(__file__).resolve().parents:
    if (_p / "repo_paths.py").is_file():
        if str(_p) not in sys.path:
            sys.path.insert(0, str(_p))
        break
else:
    raise RuntimeError("Not inside project Python tree")
from repo_paths import repo_root

PROJECT_ROOT = repo_root()
SQL_DIR = os.path.join(PROJECT_ROOT, "sql", "budget")
DATA_DIR = os.path.join(PROJECT_ROOT, "data", "projects")
SQL_FILE = os.path.join(SQL_DIR, "projects_missing_budget_by_source_all_phases.sql")
OUT_CSV = os.path.join(PROJECT_ROOT, "data", "projects", "projects_missing_budget_by_source_all_phases.csv")
OUT_XLSX = os.path.join(PROJECT_ROOT, "data", "projects", "projects_missing_budget_by_source_all_phases.xlsx")


def get_query():
    if os.path.isfile(SQL_FILE):
        with open(SQL_FILE, "r", encoding="utf-8") as f:
            return f.read().strip()
    raise FileNotFoundError(SQL_FILE)


def main():
    from edp_connection import execute_query

    query = get_query()
    print("Running projects missing budget by source and phase (heatmap)...")
    columns, results = execute_query(query)
    print(f"Retrieved {len(results)} rows, columns: {columns}")

    os.makedirs(DATA_DIR, exist_ok=True)
    with open(OUT_CSV, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(columns)
        w.writerows(results)
    print(f"Results written to {OUT_CSV}")

    try:
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Missing budget by phase and source"
        for c, col in enumerate(columns, 1):
            ws.cell(row=1, column=c, value=col)
        for r, row in enumerate(results, 2):
            for c, val in enumerate(row, 1):
                ws.cell(row=r, column=c, value=val)
        wb.save(OUT_XLSX)
        print(f"Results written to {OUT_XLSX}")
    except ImportError:
        try:
            import pandas as pd
            pd.DataFrame(results, columns=columns).to_excel(
                OUT_XLSX, index=False, sheet_name="Missing budget by phase and source"
            )
            print(f"Results written to {OUT_XLSX} (via pandas)")
        except Exception as e:
            print(f"Skipping Excel: {e}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
