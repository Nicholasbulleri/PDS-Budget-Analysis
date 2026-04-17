#!/usr/bin/env python3
"""
Run budget Closed by category (EXCLUDING missing/zero area) WITH property-derived sector AND project_type;
apply city->state mapping OUTSIDE Databricks.
- Uses budget_core_query_closed_by_category_exclude_missing_area_with_property_sector_with_project_type.sql
- Loads mapping from JSON/Excel/CSV; writes CSV + Excel with state column.
"""
import os

import sys
from pathlib import Path

for _p in Path(__file__).resolve().parents:
    if (_p / "repo_paths.py").is_file():
        if str(_p) not in sys.path:
            sys.path.insert(0, str(_p))
        break
else:
    raise RuntimeError("Not inside project Python tree")
from repo_paths import repo_root

PROJECT_ROOT = repo_root()
SQL_DIR = os.path.join(PROJECT_ROOT, "sql", "budget")
DATA_DIR = os.path.join(PROJECT_ROOT, "data", "budget")
try:
    from dotenv import load_dotenv
    load_dotenv(os.path.join(PROJECT_ROOT, ".env"))
except ImportError:
    pass

import csv
import json
SQL_FILE = os.path.join(SQL_DIR, "budget_core_query_closed_by_category_exclude_missing_area_with_property_sector_with_project_type.sql")
OUT_CSV = os.path.join(DATA_DIR, "budget_query_results_curated_closed_by_category_with_state_using_mapping_excludes_missing_area_with_property_sector_and_project_type.csv")
OUT_XLSX = os.path.join(DATA_DIR, "budget_query_results_curated_closed_by_category_with_state_using_mapping_excludes_missing_area_with_property_sector_and_project_type.xlsx")
SECTOR_MAPPING_FILE = os.path.join(PROJECT_ROOT, "data", "sector", "sector_mapping.json")

CANONICAL_SECTOR_ORDER = [
    "Office", "Retail", "Industrial and Logistics", "Data Center", "Hotels/Hospitality",
    "Residential", "Education", "Mixed-use", "Infrastructure/Energy", "Special Purpose Facility",
]

MAPPING_NAMES = [
    "city_to_state_mapping.json",
    "city_to_state_mapping.xlsx",
    "city_to_state_mapping.xls",
    "city_to_state_mapping.csv",
]


def normalize_key(city, country):
    c = (city or "").strip().lower()
    co = (country or "").strip().lower()
    return (c, co)


# US country variants: mapping file uses "United States"; Ingenious/query may return "US" or "USA"
US_COUNTRY_ALIASES = {"us", "usa", "united states", "united states of america"}


def country_for_state_lookup(country_val):
    """Return canonical country for state lookup so (city, country) matches the mapping file.
    Mapping uses 'United States'; Ingenious/query often returns 'US' or 'USA'."""
    raw = (country_val or "").strip()
    c = raw.lower()
    if c in US_COUNTRY_ALIASES:
        return "United States"
    return raw if raw else "United States"  # blank treated as US for lookup


def load_mapping_from_csv(path):
    rows = []
    with open(path, "r", encoding="utf-8") as f:
        r = csv.DictReader(f)
        for row in r:
            rows.append({
                "city": (row.get("city") or "").strip(),
                "country": (row.get("country") or "").strip(),
                "state_abbrev": (row.get("state_abbrev") or "").strip(),
            })
    return rows


def load_mapping_from_json(path):
    with open(path, "r", encoding="utf-8") as f:
        data = json.load(f)
    if isinstance(data, list):
        rows = []
        for item in data:
            rows.append({
                "city": (item.get("city") or "").strip(),
                "country": (item.get("country") or "").strip(),
                "state_abbrev": (item.get("state_abbrev") or item.get("state") or "").strip(),
            })
        return rows
    if isinstance(data, dict):
        if "cities" in data or "mappings" in data:
            arr = data.get("cities", data.get("mappings", []))
            return [{"city": (x.get("city") or "").strip(), "country": (x.get("country") or "").strip(), "state_abbrev": (x.get("state_abbrev") or x.get("state") or "").strip()} for x in arr]
        out = []
        for k, v in data.items():
            if "|" in k:
                city, country = k.split("|", 1)
                out.append({"city": city.strip(), "country": country.strip(), "state_abbrev": (v or "").strip()})
        return out if out else []
    return []


def load_mapping_from_excel(path):
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = []
        header = None
        for r in ws.iter_rows(values_only=True):
            if header is None:
                header = [str(c).strip().lower() if c else "" for c in r]
                ci = header.index("city") if "city" in header else 0
                co = header.index("country") if "country" in header else 1
                st = next((i for i, h in enumerate(header) if "state" in h), 2)
                continue
            cells = list(r) if r else []
            city = str(cells[ci]) if ci < len(cells) and cells[ci] is not None else ""
            country = str(cells[co]) if co < len(cells) and cells[co] is not None else ""
            state = str(cells[st]) if st < len(cells) and cells[st] is not None else ""
            rows.append({"city": city.strip(), "country": country.strip(), "state_abbrev": state.strip()})
        wb.close()
        return rows
    except ImportError:
        import pandas as pd
        df = pd.read_excel(path)
        df.columns = [str(c).strip().lower() for c in df.columns]
        city_col = "city" if "city" in df.columns else df.columns[0]
        country_col = "country" if "country" in df.columns else (df.columns[1] if len(df.columns) > 1 else "")
        state_col = next((c for c in df.columns if "state" in c), df.columns[2] if len(df.columns) > 2 else "")
        return [
            {"city": str(row.get(city_col, "")).strip(), "country": str(row.get(country_col, "")).strip(), "state_abbrev": str(row.get(state_col, "")).strip()}
            for _, row in df.iterrows()
        ]


def build_lookup(rows):
    lookup = {}
    for r in rows:
        key = normalize_key(r["city"], r["country"])
        if key not in lookup:
            lookup[key] = r["state_abbrev"] or ""
    return lookup


def load_sector_mapping():
    if not os.path.isfile(SECTOR_MAPPING_FILE):
        return {}
    with open(SECTOR_MAPPING_FILE, "r", encoding="utf-8") as f:
        data = json.load(f)
    lookup = {}
    for canon in CANONICAL_SECTOR_ORDER:
        raw_list = data.get(canon)
        if not isinstance(raw_list, list):
            continue
        for raw in raw_list:
            s = (raw or "").strip()
            if s:
                lookup[s] = canon
                lookup[s.lower()] = canon
    return lookup


def find_and_load_mapping():
    for name in MAPPING_NAMES:
        path = os.path.join(DATA_DIR, name)
        if not os.path.isfile(path):
            continue
        ext = os.path.splitext(name)[1].lower()
        if ext == ".json":
            rows = load_mapping_from_json(path)
        elif ext in (".xlsx", ".xls"):
            rows = load_mapping_from_excel(path)
        else:
            rows = load_mapping_from_csv(path)
        lookup = build_lookup(rows)
        print(f"Loaded {len(rows)} mapping rows from {name} -> {len(lookup)} unique city/country keys.")
        return lookup
    raise FileNotFoundError("No mapping file found. Add one of: " + ", ".join(MAPPING_NAMES))


def get_query():
    if os.path.isfile(SQL_FILE):
        with open(SQL_FILE, "r", encoding="utf-8") as f:
            return f.read().strip()
    raise FileNotFoundError(SQL_FILE)


def main():
    from edp_connection import execute_query

    lookup = find_and_load_mapping()
    sector_lookup = load_sector_mapping()
    default_sector = "Special Purpose Facility"
    if sector_lookup:
        print(f"Loaded sector mapping: {len(sector_lookup) // 2} raw values -> 10 canonical sectors.")

    query = get_query()
    print("Running budget Closed by category (excl. missing area, property sector, project type)...")
    columns, results = execute_query(query)
    print(f"Retrieved {len(results)} rows, columns: {columns}")

    out_columns = []
    for c in columns:
        out_columns.append(c)
        if c == "city":
            out_columns.append("state")
    if "state" not in out_columns:
        out_columns.append("state")

    si = columns.index("sector") if "sector" in columns else -1
    unmapped_sectors = set()

    out_rows = []
    for row in results:
        row_list = list(row)
        ci = columns.index("city") if "city" in columns else 0
        co = columns.index("country") if "country" in columns else 1
        city_val = row_list[ci] if ci < len(row_list) else None
        country_val = row_list[co] if co < len(row_list) else None
        # Use canonical country so "US"/"USA" match mapping file's "United States" (Ingenious often uses "US")
        country_for_lookup = country_for_state_lookup(country_val)
        state_val = lookup.get(normalize_key(city_val, country_for_lookup), "")
        if state_val == "" and (country_val or "").strip().lower() not in US_COUNTRY_ALIASES:
            state_val = lookup.get(normalize_key(city_val, (country_val or "").strip()), "")

        if si >= 0 and si < len(row_list) and sector_lookup:
            raw = row_list[si]
            s = (raw or "").strip()
            if s:
                mapped = sector_lookup.get(s) or sector_lookup.get(s.lower()) or default_sector
                if mapped == default_sector and (sector_lookup.get(s) is None and sector_lookup.get(s.lower()) is None):
                    unmapped_sectors.add(s)
                row_list = list(row_list)
                row_list[si] = mapped
            elif s == "":
                row_list = list(row_list)
                row_list[si] = None

        new_row = []
        for i, c in enumerate(columns):
            new_row.append(row_list[i] if i < len(row_list) else None)
            if c == "city":
                new_row.append(state_val or None)
        out_rows.append(new_row)

    if unmapped_sectors:
        print(f"Unmapped sector values (mapped to '{default_sector}'): {sorted(unmapped_sectors)}")

    with open(OUT_CSV, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(out_columns)
        w.writerows(out_rows)
    print(f"Results written to {OUT_CSV}")

    try:
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Closed + sector + project type"[:31]
        for c, col in enumerate(out_columns, 1):
            ws.cell(row=1, column=c, value=col)
        for r, row in enumerate(out_rows, 2):
            for c, val in enumerate(row, 1):
                ws.cell(row=r, column=c, value=val)
        wb.save(OUT_XLSX)
        print(f"Results written to {OUT_XLSX}")
    except ImportError:
        import pandas as pd
        pd.DataFrame(out_rows, columns=out_columns).to_excel(OUT_XLSX, index=False, sheet_name="Closed + sector + project type"[:31])
        print(f"Results written to {OUT_XLSX} (via pandas)")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
