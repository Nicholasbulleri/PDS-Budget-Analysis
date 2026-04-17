#!/usr/bin/env python3
"""
Run budget Closed by category + state/country derived. Outputs CSV and Excel.
"""
import csv
import os

try:
    from dotenv import load_dotenv
    try:
        load_dotenv()
    except (PermissionError, FileNotFoundError):
        pass
except ImportError:
    pass

import sys
from pathlib import Path

for _p in Path(__file__).resolve().parents:
    if (_p / "repo_paths.py").is_file():
        if str(_p) not in sys.path:
            sys.path.insert(0, str(_p))
        break
else:
    raise RuntimeError("Not inside project Python tree")
from repo_paths import repo_root

PROJECT_ROOT = repo_root()
SQL_DIR = os.path.join(PROJECT_ROOT, "sql", "budget")
DATA_DIR = os.path.join(PROJECT_ROOT, "data", "budget")
SQL_FILE = os.path.join(SQL_DIR, "budget_core_query_closed_by_category_with_state.sql")
OUT_CSV = os.path.join(DATA_DIR, "budget_query_results_curated_closed_by_category_with_state.csv")
OUT_XLSX = "budget_query_results_curated_closed_by_category_with_state.xlsx"


def get_query():
    if os.path.isfile(SQL_FILE):
        with open(SQL_FILE, "r", encoding="utf-8") as f:
            return f.read().strip()
    raise FileNotFoundError(SQL_FILE)


def main():
    from edp_connection import execute_query

    query = get_query()
    print("Running budget Closed by category + state/country derived...")
    columns, results = execute_query(query)
    print(f"Retrieved {len(results)} rows, columns: {columns}")

    with open(OUT_CSV, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(columns)
        w.writerows(results)
    print(f"Results written to {OUT_CSV}")

    try:
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Closed by Category + State"
        for c, col in enumerate(columns, 1):
            ws.cell(row=1, column=c, value=col)
        for r, row in enumerate(results, 2):
            for c, val in enumerate(row, 1):
                ws.cell(row=r, column=c, value=val)
        wb.save(OUT_XLSX)
        print(f"Results written to {OUT_XLSX}")
    except ImportError:
        import pandas as pd
        pd.DataFrame(results, columns=columns).to_excel(OUT_XLSX, index=False, sheet_name="Closed by Category + State")
        print(f"Results written to {OUT_XLSX} (via pandas)")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
