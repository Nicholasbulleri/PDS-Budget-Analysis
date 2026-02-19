#!/usr/bin/env python3
"""
Rebuild taskname -> category mapping from the latest budget Closed query results
(budget_query_results_curated_closed.csv) so all tasknames in that result set are mapped.
Uses the same categorization logic as build_taskname_mapping.py.
"""
import csv
import json
import re

# Reuse categorization logic: import from build_taskname_mapping
from build_taskname_mapping import categorize

CLOSED_RESULTS_CSV = "budget_query_results_curated_closed.csv"
MAPPING_JSON = "budget_taskname_category_mapping.json"
MAPPING_CSV = "budget_taskname_category_mapping.csv"
MAPPING_XLSX = "budget_taskname_category_mapping.xlsx"
SUMMARY_TXT = "budget_mapping_summary.txt"


def get_tasknames_from_closed_results():
    """Unique tasknames from the Closed budget query CSV."""
    with open(CLOSED_RESULTS_CSV, "r", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        names = set()
        for row in reader:
            tn = row.get("taskname", "").strip()
            if tn:
                names.add(tn)
    return sorted(names)


def main():
    print(f"Reading tasknames from {CLOSED_RESULTS_CSV}...")
    tasknames = get_tasknames_from_closed_results()
    print(f"Found {len(tasknames)} unique tasknames.")

    mapping = {}
    by_cat = {"FF&E + Millwork": [], "Soft Costs": [], "Construction": [], "Uncategorized": []}
    for name in tasknames:
        cat = categorize(name)
        mapping[name] = cat
        by_cat[cat].append(name)

    out = {
        "description": "Task name to budget category mapping (from budget Closed query results). Commercial real estate project/construction management.",
        "source": "budget_query_results_curated_closed.csv (work_dynamics.curated generictask + project, Closed USD, budget + Cost Code Item)",
        "categories": {
            "FF&E + Millwork": "Furniture, finishings, equipment and millwork; decorative materials; IT/security integrator; POS; storage shelving; owner-supplied items",
            "Soft Costs": "Designer, architect, engineers, expeditor, project management, special inspections, landmark consultant",
            "Construction": "General contractor and construction trade costs",
        },
        "taskname_to_category": mapping,
    }
    with open(MAPPING_JSON, "w", encoding="utf-8") as f:
        json.dump(out, f, indent=2, ensure_ascii=False)
    print(f"Wrote {MAPPING_JSON}")

    with open(MAPPING_CSV, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f)
        w.writerow(["taskname", "category"])
        for name in tasknames:
            w.writerow([name, mapping[name]])
    print(f"Wrote {MAPPING_CSV}")

    try:
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Taskname to Category"
        ws.cell(row=1, column=1, value="taskname")
        ws.cell(row=1, column=2, value="category")
        for i, name in enumerate(tasknames, 2):
            ws.cell(row=i, column=1, value=name)
            ws.cell(row=i, column=2, value=mapping[name])
        wb.save(MAPPING_XLSX)
        print(f"Wrote {MAPPING_XLSX}")
    except ImportError:
        import pandas as pd
        pd.DataFrame({"taskname": tasknames, "category": [mapping[n] for n in tasknames]}).to_excel(
            MAPPING_XLSX, index=False, sheet_name="Taskname to Category"
        )
        print(f"Wrote {MAPPING_XLSX} (via pandas)")

    summary_lines = []
    for cat in ["FF&E + Millwork", "Soft Costs", "Construction", "Uncategorized"]:
        summary_lines.append(f"## {cat}: {len(by_cat[cat])} task names")
    with open(SUMMARY_TXT, "w", encoding="utf-8") as f:
        f.write("\n".join(summary_lines) + "\n\n")
        f.write("--- Uncategorized (review these) ---\n")
        for n in by_cat["Uncategorized"][:300]:
            f.write(n + "\n")
        if len(by_cat["Uncategorized"]) > 300:
            f.write(f"... and {len(by_cat['Uncategorized']) - 300} more\n")
    print(f"Wrote {SUMMARY_TXT}")

    print("\n" + "\n".join(summary_lines))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
